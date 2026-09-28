# -*- coding: utf-8 -*-
"""Insert the formatted Results section (text + Word tables + embedded figures
+ captions) into docs/manuscript-publication.docx, renumber sections, and make
a backup first. Run: python3 -m tools.apply_results"""
import docx, shutil, sys
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "docs/manuscript-publication.docx"
shutil.copy(SRC, "docs/manuscript-publication.backup.docx")
d = docx.Document(SRC)


def fstyle(*keys):
    for s in d.styles:
        n = s.name.lower()
        if all(k in n for k in keys):
            return s.name
    return None

H1 = fstyle("heading1") or fstyle("2.1")
H2 = fstyle("heading2") or fstyle("2.2")
BODY = fstyle("3.1_text") or fstyle("text")
FCAP = fstyle("figure_caption") or fstyle("figure")
TCAP = fstyle("table_caption") or fstyle("table", "cap")
print("styles:", H1, "|", H2, "|", BODY, "|", FCAP, "|", TCAP)

paras = d.paragraphs
def findp(pred):
    for p in paras:
        if pred(p.text.strip().lower()):
            return p
    return None

res = findp(lambda t: t.startswith("3. results"))
dis = findp(lambda t: t.startswith("4. discussion"))
con = findp(lambda t: t.startswith("5. conclusions"))
pat = findp(lambda t: t.startswith("6. patents"))
assert res and dis, "Results/Discussion headings not found"

# delete template body between Results heading and Discussion heading
body = d.element.body
el = res._p.getnext()
while el is not None and el is not dis._p:
    nxt = el.getnext()
    body.remove(el)
    el = nxt

def settext(p, t):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(t)

settext(res, "4. Results")
if dis: settext(dis, "5. Discussion")
if con: settext(con, "6. Conclusions")
if pat: settext(pat, "7. Patents")

def ins(text, style=BODY):
    p = dis.insert_paragraph_before(text)
    try:
        p.style = d.styles[style]
    except Exception:
        pass
    return p

def insfig(path):
    p = dis.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(5.0))
    return p

def instable(headers, rows, caption):
    ins(caption, TCAP)                      # caption ABOVE table (MDPI)
    t = d.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        t.style = d.styles["Table Grid"]
    except Exception:
        pass
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = h
        for r in c.paragraphs[0].runs:
            r.font.bold = True; r.font.size = Pt(9)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.rows[i + 1].cells[j]; c.text = str(v)
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(9)
    dis._p.addprevious(t._tbl)              # move table before Discussion

# ---- 4.1 ----
ins("4.1. Algorithm Comparison on the Rule-Based WAF", H2)
ins("The union corpus isolates the optimizer, since every agent shares the environment, the reward, and the 1.5×10⁵-step budget. Table 1 reports the outcome against ModSecurity. TRPO reaches IFNR = 99.1% at SPBARC = 6.07, ahead of PPO at 88.9% and A2C at 76.9%, while an unguided random mutator clears only 3.7% at a cost of 216 requests per bypass. The ranking is monotonic in both effectiveness and efficiency, so the trust-region method dominates on the metric that matters to an attacker and on the metric that matters to a defender watching request volume.")
instable(["Algorithm", "IFNR", "SPBARC", "Trivial", "Medium", "Complex"],
         [["TRPO", "99.1%", "6.07", "100%", "100%", "97.2%"],
          ["PPO", "88.9%", "6.89", "100%", "100%", "66.7%"],
          ["A2C", "76.9%", "10.08", "100%", "97.2%", "33.3%"],
          ["Random (baseline)", "3.7%", "216", "0%", "–", "–"]],
         "Table 1. Algorithm comparison on the union corpus against ModSecurity (N = 108).")
ins("The per-tier figures locate the separation. All three agents saturate the trivial and medium tiers near 100%, and the gap opens only on the complex tier, where the success rate falls to 97.2%, 66.7%, and 33.3% for TRPO, PPO, and A2C (Figure 3). The complex tier is precisely the set that demands the longest mutation chains, so the spread measures how well each optimizer sustains a sparse-reward search over a long horizon rather than how well it handles SQLi in general. This dependence on horizon, not on syntax, is the first signal that the comparison probes optimization stability.")
insfig("figures/fig_tier_sr.png")
ins("Figure 3. Per-tier success rate on the union corpus against ModSecurity. All algorithms saturate the trivial and medium tiers; the ranking separates only on the complex tier, which requires the longest mutation chains.", FCAP)

# ---- 4.2 ----
ins("4.2. Generalization to Error-Based Payloads and Stability", H2)
ins("Error-based payloads invert the difficulty, and a single run would misrepresent the result, so each algorithm is trained under three seeds. Table 2 reports the mean and standard deviation. The best IFNR collapses from 99.1% on union to 30.3% on error, which confirms that error-based extraction is the harder regime. The decisive finding is the dispersion. TRPO and PPO hold a standard deviation near 2%, whereas A2C swings across σ = 14.4%, a spread roughly 7× wider that ranges from 5.6% to 30.6% on identical settings (Figure 4). A single unlucky seed places A2C at near-zero and would license a false claim of collapse; the three-seed protocol shows instead an unstable optimizer with a competitive mean.")
instable(["Algorithm", "IFNR (mean ± σ)", "Range", "SPBARC"],
         [["TRPO", "30.3% ± 2.3%", "27.8–32.4%", "28–45"],
          ["PPO", "26.3% ± 1.9%", "24.1–27.8%", "30–38"],
          ["A2C", "22.3% ± 14.4%", "5.6–30.6%", "41–269"]],
         "Table 2. Error corpus over three seeds against ModSecurity (mean ± standard deviation, N = 108).")
insfig("figures/fig_seed_variance.png")
ins("Figure 4. Induced false negative rate on the error corpus across three seeds. TRPO and PPO cluster within σ ≈ 2%, whereas A2C spans σ = 14.4% (5.6–30.6%), roughly seven times wider.", FCAP)
ins("The variance maps directly onto the algorithm design. A2C imposes no bound on the update step, so an early unfavorable batch overwrites a partially formed policy and the run never recovers, while the trust-region and clipped objectives of TRPO and PPO cap each step and reproduce across seeds. Stability, not mean performance, is therefore the property that distinguishes the methods under sparse reward. This result also carries a practical caution: any RL evaluation that reports A2C from one seed reports noise.")

# ---- 4.3 ----
ins("4.3. Mutation Ordering as a Causal Factor", H2)
ins("The last-action component of the state lets the policy condition on order, and the trained trajectories expose how much order matters. Counting operator pairs whose two orderings yield a success gap, we find 68, 146, and 137 ordering-dependent pairs for TRPO, PPO, and A2C out of roughly 1,200 pairs each. Retaining only the pairs on which at least two algorithms agree leaves 35 cross-algorithm consensus pairs, which removes any single-optimizer artifact. The agreement is the evidence: an effect that survives independent learners is a property of the environment, not of one search procedure.")
ins("The pairs encode preconditions rather than preferences (Figure 5). The ordering ident_backtick → hex_to_char succeeds in 98.0% of episodes while its reverse scores 0.0%, because the hexadecimal rewrite becomes reachable only after the identifier escape has fired. The dependency reappears with smaller but consistent gaps elsewhere: hex_to_char → case runs 93.2% against 27.4%, and null_byte → agg_swap runs 91.4% against 40.1%. The pair newline → func_sp_nbsp recurs across all three algorithms with a positive gap, which marks a shared structural dependency in the rule topology. Order is a causal variable in the search, and a set-based formulation that discards it forfeits a measurable fraction of the reachable bypasses.")
insfig("figures/fig_ordering.png")
ins("Figure 5. Forward versus reversed success rate for the top ordering-dependent mutation pairs (TRPO). Reversing the order collapses the success rate, with ident_backtick → hex_to_char falling from 98% to 0%.", FCAP)

# ---- 4.4 ----
ins("4.4. Cross-Paradigm Transfer", H2)
ins("The transfer experiment moves a fixed policy from the rule-based to the learning-based firewall. The result is categorical. A ModSecurity-trained policy scores IFNR = 0% on Safeline for all three algorithms and both corpora, with every one of more than 900 mutated requests returning a block. A zero this clean demands a control, so we verify that Safeline is bypassable in principle: a benign request returns 200, a textbook injection returns 403, and a hand-crafted payload returns 200 with extracted data. The transfer gap therefore measures policy mismatch, not an impregnable target.")
instable(["Setting", "WAF type", "WAF-evasion", "Exfiltration (IFNR)"],
         [["ModSecurity", "rule-based", "high", "76.9–99.1%"],
          ["Safeline, zero-shot transfer", "learning-based", "≈ 0%", "0%"],
          ["Safeline, trained directly", "learning-based", "≈ 89%", "0%"]],
         "Table 3. The three observed regimes. Evasion and exfiltration coincide on the rule engine and separate on the classifier.")
ins("Training directly against Safeline sharpens the picture. The agent learns to clear the firewall in about 89% of episodes, yet it exfiltrates nothing: 846 of 972 requests return status 200 paired with a SQL error, which means the payload passed the classifier only by ceasing to be valid SQL. The mutations that fool the learned boundary are the same mutations that dismantle the attack, so evasion and exfiltration, a single event on the rule engine, become mutually exclusive on the classifier. This decoupling is the structural property the transfer measures, and it indicates that defeating a learning-based firewall requires a semantics-aware operator family that the rule-oriented action space does not contain.")

out = SRC
try:
    d.save(out)
except PermissionError:
    out = "docs/manuscript-publication_fixed.docx"
    print("[!] original locked -> saved", out, file=sys.stderr)
    d.save(out)
print("saved:", out)
print("backup: docs/manuscript-publication.backup.docx")
