# -*- coding: utf-8 -*-
"""Strengthen citations: remap existing [n], insert 7 new citations at anchors,
rebuild the reference list (16 -> 23), keep ascending order. Run-level edits
preserve equations. Backup first. Run: python3 -m tools.add_refs"""
import docx, shutil, re, sys, copy
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = "docs/manuscript-publication.docx"
shutil.copy(SRC, "docs/manuscript-publication.backup.docx")
d = docx.Document(SRC)
paras = d.paragraphs
ref_i = next(i for i, p in enumerate(paras) if p.text.strip().lower() == "references")
body = paras[:ref_i]

# ---- Step A: remap existing citation numbers ----
OLDNEW = {1:1, 2:2, 3:4, 4:5, 5:6, 6:7, 7:8, 8:10, 9:11, 10:12, 11:13, 12:14,
          13:18, 14:20, 15:22, 16:23}
def remap(m):
    nums = [int(x) for x in re.split(r'\s*,\s*', m.group(1))]
    return "[" + ",".join(str(x) for x in sorted(OLDNEW.get(n, n) for n in nums)) + "]"
CITE = re.compile(r'\[(\d+(?:\s*,\s*\d+)*)\]')
for p in body:
    for r in p.runs:
        if "[" in r.text:
            r.text = CITE.sub(remap, r.text)

# ---- Step B: insert new citations at unique anchors (run-level) ----
INSERTS = [
    ("executes data as code", "executes data as code [3]"),                       # Halfond
    ("slips through [8]", "slips through [8,9]"),                                  # Goodfellow
    ("classify it [8]", "classify it [8,15]"),                                     # Rosca
    ("crosses the boundary as benign [8]", "crosses the boundary as benign [8,16,17]"),  # Szegedy, Chakraborty
    ("recasts the search as sequential control", "recasts the search as sequential control [19]"),  # Anderson
    ("finite-horizon Markov decision process", "finite-horizon Markov decision process [21]"),      # Sutton & Barto
]
done = {a: False for a, _ in INSERTS}
for p in body:
    for r in p.runs:
        for a, b in INSERTS:
            if not done[a] and a in r.text:
                r.text = r.text.replace(a, b, 1); done[a] = True
missing = [a for a, ok in done.items() if not ok]
print("inserted:", sum(done.values()), "/", len(INSERTS))
if missing:
    print("[!] anchors NOT found (add manually):")
    for a in missing: print("    -", a)

# ---- Step C: rebuild reference list (23 entries) ----
REFS = [
 "OWASP Foundation. OWASP Top 10:2021. Available online: https://owasp.org/Top10/ (accessed on 22 June 2026).",
 "MITRE. 2024 CWE Top 25 Most Dangerous Software Weaknesses. Available online: https://cwe.mitre.org/top25/ (accessed on 22 June 2026).",
 "Halfond, W.G.J.; Viegas, J.; Orso, A. A Classification of SQL-Injection Attacks and Countermeasures. In Proceedings of the IEEE International Symposium on Secure Software Engineering (ISSSE), Arlington, VA, USA, 13-15 March 2006.",
 "OWASP Foundation. OWASP ModSecurity Core Rule Set (CRS). Available online: https://coreruleset.org/ (accessed on 22 June 2026).",
 "Demetrio, L.; Valenza, A.; Costa, G.; Lagorio, G. WAF-A-MoLE: Evading Web Application Firewalls through Adversarial Machine Learning. In Proceedings of the 35th Annual ACM Symposium on Applied Computing (SAC '20), Brno, Czech Republic, 30 March-3 April 2020; pp. 1745-1752.",
 "Qu, Z.; Ling, X.; Wang, T.; Chen, X.; Ji, S.; Wu, C. AdvSQLi: Generating Adversarial SQL Injections against Real-World WAF-as-a-Service. IEEE Trans. Dependable Secur. Comput. 2024. Preprint: arXiv:2401.02615.",
 "Chaitin Tech. SafeLine Web Application Firewall. Available online: https://waf.chaitin.com/ (accessed on 22 June 2026).",
 "Guan, Y.; He, J.; Li, T.; Zhao, H.; Ma, B. SSQLi: A Black-Box Adversarial Attack Method for SQL Injection Based on Reinforcement Learning. Future Internet 2023, 15, 133.",
 "Goodfellow, I.J.; Shlens, J.; Szegedy, C. Explaining and Harnessing Adversarial Examples. In Proceedings of the 3rd International Conference on Learning Representations (ICLR), San Diego, CA, USA, 7-9 May 2015.",
 "Leung, D.; Tsai, O.; Hashemi, K.; Tayebi, B.; Tayebi, M.A. XploitSQL: Advancing Adversarial SQL Injection Attack Generation with Language Models and Reinforcement Learning. In Proceedings of the 33rd ACM International Conference on Information and Knowledge Management (CIKM '24), Boise, ID, USA, 21-25 October 2024; pp. 4653-4660.",
 "Zhang, B.; Liu, C.; Ren, R.; Wang, Q.; Ren, J. BWAFSQLi: Bypassing Web Application Firewall with Adversarial SQL Injections. ACM Trans. Softw. Eng. Methodol. 2026. https://doi.org/10.1145/3788286.",
 "Schulman, J.; Wolski, F.; Dhariwal, P.; Radford, A.; Klimov, O. Proximal Policy Optimization Algorithms. arXiv 2017, arXiv:1707.06347.",
 "Schulman, J.; Levine, S.; Abbeel, P.; Jordan, M.; Moritz, P. Trust Region Policy Optimization. In Proceedings of the 32nd International Conference on Machine Learning (ICML), Lille, France, 6-11 July 2015; pp. 1889-1897.",
 "Mnih, V.; Badia, A.P.; Mirza, M.; Graves, A.; Lillicrap, T.; Harley, T.; Silver, D.; Kavukcuoglu, K. Asynchronous Methods for Deep Reinforcement Learning. In Proceedings of the 33rd International Conference on Machine Learning (ICML), New York, NY, USA, 19-24 June 2016; pp. 1928-1937.",
 "Rosca, C.-M.; Stancu, A.; Popescu, C. Machine Learning Models for SQL Injection Detection. Electronics 2025, 14, 3420.",
 "Szegedy, C.; Zaremba, W.; Sutskever, I.; Bruna, J.; Erhan, D.; Goodfellow, I.; Fergus, R. Intriguing Properties of Neural Networks. In Proceedings of the 2nd International Conference on Learning Representations (ICLR), Banff, AB, Canada, 14-16 April 2014.",
 "Chakraborty, A.; Alam, M.; Dey, V.; Chattopadhyay, A.; Mukhopadhyay, D. A Survey on Adversarial Attacks and Defences. CAAI Trans. Intell. Technol. 2021, 6, 25-45.",
 "Appelt, D.; Nguyen, C.D.; Panichella, A.; Briand, L.C. A Machine-Learning-Driven Evolutionary Approach for Testing Web Application Firewalls. IEEE Trans. Reliab. 2018, 67, 733-757.",
 "Anderson, H.S.; Kharkar, A.; Filar, B.; Evans, D.; Roth, P. Learning to Evade Static PE Machine Learning Malware Models via Reinforcement Learning. arXiv 2018, arXiv:1801.08917.",
 "Hemmati, M.; Hadavi, M.A. Bypassing Web Application Firewalls Using Deep Reinforcement Learning. ISC Int. J. Inf. Secur. 2022, 14, 131-145. https://doi.org/10.22042/isecure.2022.323140.744.",
 "Sutton, R.S.; Barto, A.G. Reinforcement Learning: An Introduction, 2nd ed.; MIT Press: Cambridge, MA, USA, 2018.",
 "Ng, A.Y.; Harada, D.; Russell, S. Policy Invariance Under Reward Transformations: Theory and Application to Reward Shaping. In Proceedings of the 16th International Conference on Machine Learning (ICML), Bled, Slovenia, 27-30 June 1999; pp. 278-287.",
 "Dhakal, A. (Audi-1). sqli-labs: Test Platform for SQL Injection. Available online: https://github.com/Audi-1/sqli-labs (accessed on 22 June 2026).",
]
refp = [p for p in paras[ref_i+1:] if (p.style.name or "").startswith("MDPI_8.1") and p.text.strip()]
def settext(p, t):
    for r in list(p.runs): r._element.getparent().remove(r._element)
    p.add_run(t)
def setnum(p):
    pPr = p._p.get_or_add_pPr()
    numPr = pPr.find(qn('w:numPr'))
    if numPr is None:
        numPr = pPr.makeelement(qn('w:numPr'), {}); pPr.append(numPr)
    for tag in ('w:ilvl','w:numId'):
        e = numPr.find(qn(tag))
        if e is None:
            e = numPr.makeelement(qn(tag), {}); numPr.append(e)
    numPr.find(qn('w:ilvl')).set(qn('w:val'),'0')
    numPr.find(qn('w:numId')).set(qn('w:val'),'11')

prev = None
for i, txt in enumerate(REFS):
    if i < len(refp):
        p = refp[i]; settext(p, txt); setnum(p); prev = p
    else:
        new_el = copy.deepcopy(prev._p); prev._p.addnext(new_el)
        p = Paragraph(new_el, prev._parent); settext(p, txt); setnum(p); prev = p
# delete leftover old ref paragraphs if existing > 23 (not the case here)
for p in refp[len(REFS):]:
    p._element.getparent().remove(p._element)

out = SRC
try:
    d.save(out)
except PermissionError:
    out = "docs/manuscript-publication_fixed.docx"
    print("[!] locked -> " + out, file=sys.stderr); d.save(out)
print("saved:", out)
