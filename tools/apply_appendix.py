# -*- coding: utf-8 -*-
"""Append 'Appendix A' with the full 51-operator table (#16). Backup first.
Run: python3 -m tools.apply_appendix"""
import docx, shutil, sys
from docx.shared import Pt
from docx.enum.text import WD_COLOR_INDEX
from seqsqli.core.mutations import MUTATIONS

SRC = "docs/manuscript-publication.docx"
shutil.copy(SRC, "docs/manuscript-publication.backup.docx")
d = docx.Document(SRC)

def fstyle(*k):
    for s in d.styles:
        n = s.name.lower()
        if all(x in n for x in k):
            return s.name
H1  = fstyle("heading1") or fstyle("2.1")
BODY = fstyle("3.1_text") or fstyle("text")
TCAP = fstyle("table_caption") or fstyle("table", "cap")

# family assignment (name -> family), div_break excluded (not in the 51 evaluated)
FAM = {
 # separator substitution
 "comment":1,"newline":1,"tab_space":1,"crlf":1,"vtab":1,"formfeed":1,"between_space":1,
 "plus_sep":1,"double_encode":1,"dbl_url_comment":1,
 # keyword obfuscation
 "case":2,"case_split":2,"url_encode":2,"url_mid_char":2,"keyword_split":2,"char_split":2,
 "ver_50000":2,"ver_00000":2,"ver_12345":2,"ver_urlenc":2,"ver_newline":2,"nested_union":2,
 "backtick":2,"double_or":2,"double_and":2,"nested_or":2,"nested_and":2,
 # value and string encoding
 "hex_encode":3,"scientific":3,"concat_char":3,"hex_to_char":3,
 # function-form variation
 "func_sp_lit":4,"func_sp_nl":4,"func_sp_tab":4,"func_sp_ff":4,"func_sp_nbsp":4,"func_sp_cmt":4,
 "func_swap_err":4,"agg_swap":4,"unhex_wrap":4,"convert_charset":4,
 # structural transformation
 "paren_space":5,"paren_full":5,"distinct":5,"param_pollute":5,"hash_nl":5,"hash_nl_all":5,"ident_backtick":5,
 # parser-divergence tricks
 "null_byte":6,"dot_prefix":6,"gbk_bypass":6,
}
FAMNAME = {1:"Separator substitution",2:"Keyword obfuscation",3:"Value and string encoding",
           4:"Function-form variation",5:"Structural transformation",6:"Parser-divergence"}

def desc(fn):
    doc = (fn.__doc__ or "").strip().splitlines()
    line = doc[0].strip() if doc else ""
    return (line[:95] + "…") if len(line) > 96 else line

rows = []
for name, fn in MUTATIONS.items():
    if name == "div_break":
        continue
    rows.append((FAM.get(name, 9), name, FAMNAME.get(FAM.get(name, 9), "?"), desc(fn)))
rows.sort(key=lambda x: (x[0], x[1]))
assert len(rows) == 51, f"expected 51 operators, got {len(rows)}"

# insertion anchor = References heading
ref = next(p for p in d.paragraphs if p.text.strip().lower() == "references")
def ins(text, style, red=False):
    p = ref.insert_paragraph_before(text)
    try: p.style = d.styles[style]
    except Exception: pass
    if red:
        for r in p.runs: r.font.highlight_color = WD_COLOR_INDEX.RED
    return p

ins("Appendix A", H1, red=True)
ins("Table A1 lists the complete action set of 51 mutation operators, grouped by the six families "
    "referenced in Section 3.2. Each operator is a pure function that rewrites the surface form of a "
    "payload while preserving its attack semantics.", BODY)
ins("Table A1. The complete set of 51 mutation operators grouped by family.", TCAP, red=True)

t = d.add_table(rows=1 + len(rows), cols=4)
try: t.style = d.styles["Table Grid"]
except Exception: pass
hdr = ["#", "Operator", "Family", "Description"]
for j, h in enumerate(hdr):
    c = t.rows[0].cells[j]; c.text = h
    for r in c.paragraphs[0].runs: r.font.bold = True; r.font.size = Pt(8)
for i, (fnum, name, fam, ds) in enumerate(rows, 1):
    vals = [str(i), name, fam, ds]
    for j, v in enumerate(vals):
        c = t.rows[i].cells[j]; c.text = v
        for r in c.paragraphs[0].runs: r.font.size = Pt(8)
ref._p.addprevious(t._tbl)   # move table before References

out = SRC
try:
    d.save(out)
except PermissionError:
    out = "docs/manuscript-publication_fixed.docx"
    print("[!] locked -> " + out, file=sys.stderr); d.save(out)
print(f"inserted Appendix A with {len(rows)} operators. saved: {out}")
