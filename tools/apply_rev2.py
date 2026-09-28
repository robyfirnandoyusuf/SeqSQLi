# -*- coding: utf-8 -*-
"""Round-2 reviewer copy-edits. Segment-rebuild for equation-free paragraphs,
run-level for paragraphs with equations/existing highlights. Red-highlight each
change. Run: python3 -m tools.apply_rev2"""
import docx, copy, sys
from docx.enum.text import WD_COLOR_INDEX
from docx.text.run import Run
from docx.oxml.ns import qn

SRC = "docs/manuscript-publication.docx"
d = docx.Document(SRC)
ps = d.paragraphs
RED = WD_COLOR_INDEX.RED

def has_math(p): return p._p.find('.//' + qn('m:oMath')) is not None

def seg_rebuild(idx, edits, keep_label=None):
    """Rebuild a (equation-free) paragraph applying edits; new text highlighted red.
    keep_label: if set, preserve a leading bold label (e.g. 'Abstract: ')."""
    p = ps[idx]
    assert not has_math(p), f"para {idx} has equation; use run-level"
    joined = "".join(r.text for r in p.runs)
    joined = joined.replace("ﬁ", "fi").replace("ﬂ", "fl")  # de-ligature
    label = ""
    if keep_label and joined.startswith(keep_label):
        label = keep_label; joined = joined[len(keep_label):]
    segs = [(joined, False)]
    misses = []
    for old, new in edits:
        done = False; out = []
        for text, hl in segs:
            if not done and not hl and old in text:
                i = text.index(old)
                if text[:i]: out.append((text[:i], False))
                out.append((new, True))
                if text[i+len(old):]: out.append((text[i+len(old):], False))
                done = True
            else:
                out.append((text, hl))
        segs = out
        if not done: misses.append(old)
    for r in list(p.runs): r._element.getparent().remove(r._element)
    if label:
        lr = p.add_run(label); lr.font.bold = True
    for text, hl in segs:
        if not text: continue
        r = p.add_run(text)
        if hl: r.font.highlight_color = RED
    return misses

def run_edit(idx, old, new, hl=None, keep_red=False):
    """Replace within the single run that holds `old` (preserves equations)."""
    p = ps[idx]
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            if hl and not keep_red:
                _split_hl(p, hl)
            return True
    return False

def _split_hl(p, target):
    for r in list(p.runs):
        i = r.text.find(target)
        if i < 0: continue
        before, after = r.text[:i], r.text[i+len(target):]
        rEl = r._r; midEl = copy.deepcopy(rEl); aftEl = copy.deepcopy(rEl)
        rEl.addnext(aftEl); rEl.addnext(midEl); r.text = before
        m = Run(midEl, p._parent); m.text = target; m.font.highlight_color = RED
        Run(aftEl, p._parent).text = after
        return True
    return False

# ---------- Para 7: Abstract (#2,#3,#4) ----------
m7 = seg_rebuild(7, [
 ("and successfully returns the extracted data",
  "and the injected query returns the extracted data"),
 ("Reversing the order of 35 mutation pairs consistently selected by all three algorithms",
  "Reversing the order of 35 mutation pairs consistently selected across at least two of the three algorithms"),
 (". This performance exceeds that of PPO at 88.9% and A2C at 76.9%.",
  ", outperforming PPO at 88.9% and A2C at 76.9%."),
], keep_label="Abstract: ")

# ---------- Para 12: Intro (#7-#17) ----------
m12 = seg_rebuild(12, [
 ("rulebased or learning-based", "rule-based or learning-based"),
 ("regular-expression rules [4]. Which makes the system easily inspectable and not too expensive to deploy.",
  "regular-expression rules [4], which makes the system easily inspectable and not too expensive to deploy."),
 ("An attacker can embed comments, encode characters or a keyword/function with another equivalent that yields the same result.",
  "An attacker can embed comments, encode characters, or substitute a keyword or function with an equivalent construct that yields the same result."),
 ("these small changes may be sufficient to avoid a payload from matching a static rule.",
  "these small changes may be sufficient to prevent a payload from matching a static rule."),
 ("This approach differs from learning based WAFs.", "This approach differs from learning-based WAFs."),
 ("which works on the basis of a classifier trained by traffic data rather than based on hand-written rules.",
  "which operates using a classifier trained on traffic data rather than hand-written rules."),
 ("This can enhance detection when the attack follows a similar logic to a known signature.",
  "This can enhance detection when the attack resembles patterns seen during training."),
 ("However, the model relies on trends it learned during training. If, however, there are significant changes to the structure or token sequence of a malicious request, then it may no longer resemble the training examples seen by the classifier.",
  "However, the model relies on trends it learned during training. If significant changes occur in the structure or token sequence of a malicious request, it may no longer resemble the training examples seen by the classifier."),
 ("the attack will be misclassified by a benign one although its inner SQL injection behavior remains unchanged",
  "the attack will be misclassified as benign, although its inner SQL injection behavior remains unchanged"),
 ("That gives rise to a version of that same weakness in both approaches.",
  "This surface-form variability is thus a shared weakness of both detection paradigms."),
 ("it is sufficient to change the representation of a currently existing payload.",
  "it is sufficient to change the representation of an existing payload."),
])

# ---------- Para 13: Intro/RW (#18-#23) ----------
m13 = seg_rebuild(13, [
 ("This exposure is measured by sending live payloads to the block or allow verdict [4, 5].",
  "This exposure is measured by sending live payloads and observing the resulting block-or-allow verdict [4, 5]."),
 ("The first experiments have an over automated search direction; for example,",
  "Early work pursued a largely automated search-based direction; for example,"),
 ("WAF-A-MoLE traverses a mutation tree walking toward the lowest possible result of the classifier [5]",
  "WAF-A-MoLE traverses a mutation tree, moving toward the payload that minimizes the classifier's confidence score [5]"),
 ("AdvSQLi extends their payload on grammar by merging it with Monte Carlo tree search [6]",
  "AdvSQLi extends payload generation with a context-free grammar combined with Monte Carlo Tree Search [6]"),
 ("The same loop is repeated verbatim as sequential control where the reward is the firewall verdict;",
  "This same search loop is reframed as a sequential control problem, in which the reward signal is the firewall verdict;"),
 ("BWAFSQLi generates grammar by convergence and adaptive operator selection in a way that the performance on its benchmark is totally independent of reinforcement learning along eleven firewalls [11].",
  "BWAFSQLi uses a grammar-based approach with adaptive operator selection, achieving strong performance across eleven firewalls without relying on reinforcement learning [11]."),
])

# ---------- Para 69: comma splice (#33) ----------
m69 = seg_rebuild(69, [
 ("both effectiveness and efficiency, TRPO achieved the highest",
  "both effectiveness and efficiency; TRPO achieved the highest"),
])

# ---------- Para 83: rounding (#46) ----------
m83 = seg_rebuild(83, [
 ("out of roughly 1,200 pairs each", "out of the 1,275 possible pairs"),
])

# ---------- Para 96: Discussion (#39) ----------
m96 = seg_rebuild(96, [
 ("The 35 mutation-order pairs consistent across all algorithms show",
  "The 35 mutation-order pairs supported by at least two of the three algorithms show"),
])

# ---------- run-level (equations / existing highlight) ----------
r46 = run_edit(46, "An episodes runs", "An episode runs", hl="An episode runs")   # #28
r53 = run_edit(53, "the failure mode then we", "the failure mode that we", hl="failure mode that we")  # #31
r56 = run_edit(56, "WAF deployments", "WAF Deployments", hl="Deployments")        # #43
r88 = run_edit(88, "bypassable in principle a textbook", "bypassable in principle: a textbook",
                keep_red=True)                                                    # #37

for tag, m in [("P7",m7),("P12",m12),("P13",m13),("P69",m69),("P83",m83),("P96",m96)]:
    print(f"{tag}: {'OK' if not m else 'MISS '+str(m)}")
print("P46(#28):", r46, "| P53(#31):", r53, "| P56(#43):", r56, "| P88(#37):", r88)

try:
    d.save(SRC)
except PermissionError:
    d.save("docs/manuscript-publication_fixed.docx"); print("[!] locked -> _fixed", file=sys.stderr)
print("saved.")
